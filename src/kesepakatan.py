"""Baca 3 file kesepakatan (format matriks) -> satu tabel panjang.

Setiap baris hasil = satu (pelanggan, produk, bulan) dengan disc-nya. Bentuk
panjang ini yang nanti dipakai mesin pencocokan: cari baris dengan PT yang sama,
produk yang sama, dan bulan yang beririsan dengan periode promo klaim.

Tabel yang dibaca:
  * Promo_Olaif_2026_1.xlsx, sheet "OLAIF " (ada spasi di akhir):
      - tabel utama (nama PT lengkap, kolom bulan Jan-Des 2026)
      - 3 tabel tambahan di bawahnya: Lotte (Apr & Jun), GO Empo AEON
        (29 Mei-10 Jun 2026), serta AEON Festival Citylink & AEON Pakuwon
        Surabaya (diskon 25%, Des 2026 - 6 Jan 2027).
  * Olaif_Q2_2025.xlsx, sheet "OLAIF (2)" (Jan-Jun 2025, nama toko).
  * Promo_Olaif_Q4_2025_1.xlsx, sheet "Sheet1" (Okt-Des 2025, nama toko).
"""
from __future__ import annotations

import calendar
import datetime as dt
import re
from pathlib import Path

import openpyxl
import pandas as pd

from src import config
from src.mapping import is_unmapped, map_customer

COLUMNS = [
    "sumber_file", "tabel", "jenis_promo", "nama_di_file", "pt", "toko",
    "produk", "cbp", "periode_mulai", "periode_akhir",
    "disc_pct", "disc_rp_pcs", "keterangan",
]

_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "mei": 5, "jun": 6,
    "juni": 6, "jul": 7, "juli": 7, "aug": 8, "agu": 8, "agus": 8, "sep": 9,
    "oct": 10, "okt": 10, "nov": 11, "dec": 12, "des": 12,
}


def _month_bounds(year: int, month: int) -> tuple[dt.date, dt.date]:
    last = calendar.monthrange(year, month)[1]
    return dt.date(year, month, 1), dt.date(year, month, last)


def _parse_disc(value) -> tuple[float | None, float | None]:
    """Ubah isi sel menjadi (disc_pct, disc_rp_pcs).

    - angka 0 < v <= 1  -> disc persen (0.1 = 10%)
    - angka > 1         -> rupiah per pcs
    - teks "13.000/ pc" -> 13000 rupiah per pcs
    """
    if value is None:
        return None, None
    if isinstance(value, bool):
        return None, None
    if isinstance(value, (int, float)):
        v = float(value)
        if v == 0:
            return None, None
        if v <= 1:
            return v, None
        return None, v
    s = str(value).strip()
    if not s:
        return None, None
    low = s.lower()
    if "/" in low or "pc" in low or "rp" in low:
        head = re.split(r"[/]", s)[0]
        digits = re.sub(r"[^\d]", "", head)
        if digits:
            return None, float(digits)
    return None, None


def _parse_period_text(s: str, default_end_year: int = 2026) -> tuple[dt.date | None, dt.date | None]:
    """Parse teks periode seperti "18 Dec - 6 Jan 2027" atau "29 May-10 Juni".

    Tahun diambil dari teks bila ada; dipakai untuk tanggal AKHIR. Bila bulan
    mulai > bulan akhir (mis. Des -> Jan), tahun mulai = tahun akhir - 1.
    """
    s = str(s)
    pairs = re.findall(r"(\d{1,2})\s*([A-Za-z]+)", s)
    if len(pairs) < 2:
        return None, None
    year_m = re.search(r"(20\d{2})", s)
    end_year = int(year_m.group(1)) if year_m else default_end_year
    (d1, m1s), (d2, m2s) = pairs[0], pairs[1]
    m1 = _MONTHS.get(m1s[:4].lower()) or _MONTHS.get(m1s[:3].lower())
    m2 = _MONTHS.get(m2s[:4].lower()) or _MONTHS.get(m2s[:3].lower())
    if not m1 or not m2:
        return None, None
    start_year = end_year if m1 <= m2 else end_year - 1
    try:
        return dt.date(start_year, m1, int(d1)), dt.date(end_year, m2, int(d2))
    except ValueError:
        return None, None


def _row(fname, tabel, jenis, nama, produk, cbp, ms, me, dpct, drp, ket=None) -> dict:
    pt, toko = map_customer(nama)
    return dict(
        sumber_file=fname, tabel=tabel, jenis_promo=jenis,
        nama_di_file=str(nama).strip(), pt=pt, toko=toko, produk=str(produk).strip(),
        cbp=float(cbp) if isinstance(cbp, (int, float)) and not isinstance(cbp, bool) else None,
        periode_mulai=ms, periode_akhir=me, disc_pct=dpct, disc_rp_pcs=drp,
        keterangan=ket,
    )


# ---------------------------------------------------------------------------
# File 2026
# ---------------------------------------------------------------------------
def _load_2026(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = "OLAIF " if "OLAIF " in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    fname = path.name
    rows: list[dict] = []

    # Kolom bulan E(5)..P(16): header berisi tanggal awal bulan.
    month_cols: dict[int, tuple[int, int]] = {}
    for c in range(5, 17):
        h = ws.cell(row=1, column=c).value
        if isinstance(h, dt.datetime):
            month_cols[c] = (h.year, h.month)

    # Tabel utama: baris 3..234 (tabel tambahan mulai baris 235).
    for r in range(3, 235):
        nama = ws.cell(row=r, column=2).value
        produk = ws.cell(row=r, column=3).value
        if not (isinstance(nama, str) and nama.strip()):
            continue
        if nama.strip() == "Nama Pelanggan":   # header tabel tambahan
            break
        if not (isinstance(produk, str) and produk.strip()):
            continue
        cbp = ws.cell(row=r, column=4).value
        for c, (yy, mm) in month_cols.items():
            dpct, drp = _parse_disc(ws.cell(row=r, column=c).value)
            if dpct is None and drp is None:
                continue
            ms, me = _month_bounds(yy, mm)
            rows.append(_row(fname, "Utama", "Reguler bulanan", nama, produk, cbp, ms, me, dpct, drp))

    rows += _load_2026_extra(ws, fname)
    return rows


def _load_2026_extra(ws, fname: str) -> list[dict]:
    rows: list[dict] = []

    # (1) Tabel tambahan Lotte: header baris 235, data 236..240.
    #     Kolom H(8) = April 2026; kolom J(10) berlabel "Apr 2026" tetapi
    #     posisinya di kolom Juni -> dianggap Juni 2026 (sesuai CLAUDE.md).
    if str(ws.cell(row=235, column=2).value or "").strip() == "Nama Pelanggan":
        for r in range(236, 241):
            nama = ws.cell(row=r, column=2).value
            produk = ws.cell(row=r, column=3).value
            if not (isinstance(nama, str) and nama.strip()):
                continue
            if not (isinstance(produk, str) and produk.strip()):
                continue
            for col, (yy, mm) in [(8, (2026, 4)), (10, (2026, 6))]:
                dpct, drp = _parse_disc(ws.cell(row=r, column=col).value)
                if dpct is None and drp is None:
                    continue
                ms, me = _month_bounds(yy, mm)
                ket = ("Tabel tambahan Lotte; kolom Juni aslinya berlabel 'Apr 2026'."
                       if mm == 6 else "Tabel tambahan Lotte.")
                rows.append(_row(fname, "Tambahan Lotte", "Promo tambahan",
                                 nama, produk, None, ms, me, dpct, drp, ket))

    # (2) GO Empo AEON: penanda "GO Empo", baris periode persis di bawahnya,
    #     lalu baris produk. Periode dibaca dari teks (29 Mei-10 Jun 2026).
    for rr in range(240, 250):
        if str(ws.cell(row=rr, column=5).value or "").strip().lower() == "go empo":
            ms, me = _parse_period_text(ws.cell(row=rr + 1, column=5).value, default_end_year=2026)
            for r in range(rr + 2, rr + 2 + 10):
                nama = ws.cell(row=r, column=2).value
                produk = ws.cell(row=r, column=3).value
                if not (isinstance(nama, str) and nama.strip()):
                    break
                if nama.strip() == "Nama Pelanggan":
                    break
                dpct, drp = _parse_disc(ws.cell(row=r, column=5).value)
                if dpct is None and drp is None:
                    continue
                rows.append(_row(fname, "GO Empo", "GO Empo", nama, produk, None,
                                 ms, me, dpct, drp, "Promo GO Empo AEON."))
            break

    # (3) AEON Festival Citylink & AEON Pakuwon Surabaya: diskon 25%,
    #     periode per baris di kolom M. Header punya "Promo" di kolom D.
    for hr in range(249, 256):
        if str(ws.cell(row=hr, column=4).value or "").strip().lower() == "promo":
            for r in range(hr + 1, hr + 1 + 15):
                nama = ws.cell(row=r, column=2).value
                produk = ws.cell(row=r, column=3).value
                if not (isinstance(nama, str) and nama.strip()):
                    continue
                if not (isinstance(produk, str) and produk.strip()):
                    continue
                promo = str(ws.cell(row=r, column=4).value or "")
                m = re.search(r"(\d{1,3})\s*%", promo)
                dpct = int(m.group(1)) / 100 if m else None
                ms, me = _parse_period_text(ws.cell(row=r, column=13).value, default_end_year=2027)
                if dpct is None:
                    continue
                rows.append(_row(fname, "AEON Festival/Pakuwon", "Diskon event",
                                 nama, produk, None, ms, me, dpct, None,
                                 f"{promo.strip()}; {str(ws.cell(row=r, column=13).value or '').strip()}"))
            break

    return rows


# ---------------------------------------------------------------------------
# File 2025 (format sederhana: Account | Description | bulan...)
# ---------------------------------------------------------------------------
def _load_simple_2025(path: Path, sheet: str, month_map: dict[int, int],
                      year: int, tabel: str, ket_col: int | None) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    fname = path.name
    rows: list[dict] = []
    for r in range(3, ws.max_row + 1):
        acc = ws.cell(row=r, column=1).value
        produk = ws.cell(row=r, column=2).value
        if not (isinstance(acc, str) and acc.strip()):
            continue
        if not (isinstance(produk, str) and produk.strip()):
            continue
        ket = ws.cell(row=r, column=ket_col).value if ket_col else None
        ket = str(ket).strip() if isinstance(ket, str) and ket.strip() else None
        for col, mm in month_map.items():
            dpct, drp = _parse_disc(ws.cell(row=r, column=col).value)
            if dpct is None and drp is None:
                continue
            ms, me = _month_bounds(year, mm)
            rows.append(_row(fname, tabel, "Reguler bulanan", acc, produk, None,
                             ms, me, dpct, drp, ket))
    return rows


# ---------------------------------------------------------------------------
# API utama
# ---------------------------------------------------------------------------
def load_kesepakatan(file_2026=None, file_q2_2025=None, file_q4_2025=None) -> pd.DataFrame:
    """Gabungkan 3 file kesepakatan menjadi satu DataFrame panjang."""
    f2026 = Path(file_2026) if file_2026 else config.FILE_2026
    fq2 = Path(file_q2_2025) if file_q2_2025 else config.FILE_Q2_2025
    fq4 = Path(file_q4_2025) if file_q4_2025 else config.FILE_Q4_2025

    rows: list[dict] = []
    if f2026.exists():
        rows += _load_2026(f2026)
    if fq2.exists():
        rows += _load_simple_2025(fq2, "OLAIF (2)", {3: 1, 4: 2, 5: 3, 6: 4, 7: 5, 8: 6},
                                  2025, "Sheet OLAIF (2) (Jan-Jun 2025)", ket_col=9)
    if fq4.exists():
        rows += _load_simple_2025(fq4, "Sheet1", {3: 10, 4: 11, 5: 12},
                                  2025, "Sheet Sheet1 (Okt-Des 2025)", ket_col=None)

    df = pd.DataFrame(rows, columns=COLUMNS)
    df["periode_mulai"] = pd.to_datetime(df["periode_mulai"])
    df["periode_akhir"] = pd.to_datetime(df["periode_akhir"])
    return df


def ringkas(df: pd.DataFrame) -> str:
    """Ringkasan teks untuk ditampilkan ke pengguna."""
    lines: list[str] = []
    lines.append(f"Total baris kesepakatan (per pelanggan x produk x bulan): {len(df):,}")
    lines.append("")
    lines.append("Per file sumber:")
    for src, n in df["sumber_file"].value_counts().sort_index().items():
        lines.append(f"  - {src}: {n:,} baris")
    lines.append("")
    lines.append("Per jenis promo:")
    for jp, n in df["jenis_promo"].value_counts().items():
        lines.append(f"  - {jp}: {n:,} baris")
    lines.append("")
    lines.append("Baris per PT (10 teratas):")
    for pt, n in df["pt"].value_counts().head(10).items():
        lines.append(f"  - {pt}: {n:,}")
    n_rp = int(df["disc_rp_pcs"].notna().sum())
    lines.append("")
    lines.append(f"Isian disc dalam Rp/pcs: {n_rp} baris")
    unmapped = sorted({n for n in df["nama_di_file"].unique() if is_unmapped(n)})
    lines.append(f"Nama pelanggan belum terpetakan: {len(unmapped)}"
                 + (f" -> {unmapped}" if unmapped else ""))
    rng = f"{df['periode_mulai'].min():%d %b %Y} s.d. {df['periode_akhir'].max():%d %b %Y}"
    lines.append(f"Rentang periode: {rng}")
    return "\n".join(lines)


if __name__ == "__main__":
    data = load_kesepakatan()
    print(ringkas(data))
