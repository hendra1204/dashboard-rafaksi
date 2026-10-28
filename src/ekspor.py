"""Ekspor hasil pengecekan + keputusan finance ke Excel (mirip contoh v5).

Dipakai tombol "Unduh Excel" di dashboard. Menghitung juga berapa yang
dibayar / ditahan / ditolak berdasarkan keputusan finance atas tiap temuan.
"""
from __future__ import annotations

import datetime as dt
import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.pencocokan import KONFIRMASI, LUAR, SESUAI
from src.referensi import fmt_rp

HEAD = PatternFill("solid", fgColor="1F4E78")
GOOD = PatternFill("solid", fgColor="E2EFDA")
WARN = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="F8CBAD")

_LABEL = {"terima": "Terima", "tahan": "Tahan", "tolak": "Tolak", "": "Belum diputuskan"}


def hitung_pembayaran(data: dict, decisions: dict) -> dict:
    """Berdasarkan keputusan atas temuan, hitung dibayar/ditahan/ditolak."""
    sesuai = data["ringkasan"]["sesuai"]
    ditagih = data["ringkasan"]["ditagih"]
    terima = tahan = tolak = 0
    for t in data["temuan"]:
        st = (decisions.get(t["id"]) or {}).get("status", "")
        if st == "terima":
            terima += t["nilai"]
        elif st == "tolak":
            tolak += t["nilai"]
        else:
            tahan += t["nilai"]       # tahan atau belum diputuskan
    return dict(ditagih=ditagih, sesuai=sesuai,
                dibayar=sesuai + terima, ditahan=tahan, ditolak=tolak,
                tambahan_terima=terima)


def _hdr(ws, headers, widths):
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def bangun_workbook(data: dict, decisions: dict, batch_label: str) -> openpyxl.Workbook:
    hasil = data["hasil"]
    pay = hitung_pembayaran(data, decisions)
    wb = openpyxl.Workbook()

    # --- Ringkasan ---
    ws = wb.active
    ws.title = "Ringkasan"
    ws.cell(1, 1, "Rekap Klaim Rafaksi - " + batch_label).font = Font(bold=True, size=13)
    ws.cell(2, 1, f"Diekspor {dt.datetime.now():%d %b %Y %H:%M}").font = Font(italic=True, color="808080")
    baris = [
        (4, "Total ditagih Filia", pay["ditagih"], None),
        (5, "Sesuai kesepakatan", data["ringkasan"]["sesuai"], GOOD),
        (6, "Perlu konfirmasi", data["ringkasan"]["konfirmasi"], WARN),
        (7, "Di luar kesepakatan", data["ringkasan"]["luar"], BAD),
        (9, "Keputusan finance", None, None),
        (10, "Layak dibayar (sesuai + temuan diterima)", pay["dibayar"], GOOD),
        (11, "Ditahan (belum diputuskan / tahan)", pay["ditahan"], WARN),
        (12, "Ditolak", pay["ditolak"], BAD),
    ]
    for r, lab, val, fill in baris:
        ws.cell(r, 1, lab)
        if lab in ("Keputusan finance",):
            ws.cell(r, 1).font = Font(bold=True)
        if val is not None:
            cell = ws.cell(r, 2, val)
            cell.number_format = "#,##0"
        if fill:
            ws.cell(r, 1).fill = fill
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16

    # --- Keputusan Finance (temuan) ---
    ws = wb.create_sheet("Keputusan Finance")
    _hdr(ws, ["Temuan", "Konteks", "Nilai (Rp)", "Keputusan", "Catatan"],
         [40, 36, 14, 16, 48])
    for i, t in enumerate(data["temuan"], 2):
        dec = decisions.get(t["id"]) or {}
        st = dec.get("status", "")
        vals = [t["judul"], t["konteks"], t["nilai"], _LABEL.get(st, st), dec.get("catatan", "")]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(i, j, v)
            if j == 3:
                cell.number_format = "#,##0"
        ws.cell(i, 4).fill = {"terima": GOOD, "tolak": BAD, "tahan": WARN}.get(st, WARN)

    # --- Per Klaim ---
    ws = wb.create_sheet("Per Klaim")
    _hdr(ws, ["No Invoice", "PT", "Toko", "Keterangan Filia", "Program",
              "Periode Mulai", "Periode Akhir", "Ditagih (Rp)", "Sesuai (Rp)",
              "Perlu Konfirmasi (Rp)", "Di Luar (Rp)", "Status", "Alasan"],
         [13, 26, 18, 32, 22, 13, 13, 13, 13, 16, 12, 20, 44])
    for i, h in enumerate(hasil, 2):
        c = h.claim
        vals = [c.invoice, c.pt, c.toko, c.filia_label, c.program,
                c.periode_mulai, c.periode_akhir, round(c.filia_amount),
                round(h.sesuai), round(h.konfirmasi), round(h.luar), h.status, h.alasan]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(i, j, v)
            if j in (6, 7):
                cell.number_format = "dd mmm yyyy"
            if j in (8, 9, 10, 11):
                cell.number_format = "#,##0"
        ws.cell(i, 12).fill = (GOOD if h.status == SESUAI else BAD if h.status == LUAR else WARN)

    # --- Per Supermarket ---
    ws = wb.create_sheet("Per Supermarket")
    _hdr(ws, ["Supermarket", "PT", "Klaim", "Ditagih (Rp)", "Sesuai (Rp)",
              "Perlu Konfirmasi (Rp)", "Di Luar (Rp)", "Komposisi"],
         [22, 28, 8, 14, 13, 16, 12, 40])
    for i, p in enumerate(data["per_supermarket"], 2):
        for j, v in enumerate([p["toko"], p["pt"], p["klaim"], p["ditagih"],
                               p["sesuai"], p["konf"], p["luar"], p["komposisi"]], 1):
            cell = ws.cell(i, j, v)
            if j in (4, 5, 6, 7):
                cell.number_format = "#,##0"

    # --- Di Luar Kesepakatan ---
    ws = wb.create_sheet("Di Luar Kesepakatan")
    _hdr(ws, ["No Invoice", "PT", "Toko/Store", "Produk (lampiran)", "Nilai (Rp)", "Catatan"],
         [13, 24, 14, 42, 13, 40])
    i = 2
    for h in hasil:
        for d in h.rincian:
            if d["status"] == LUAR:
                for j, v in enumerate([h.claim.invoice, h.claim.pt, d["store"],
                                       d["product"], round(d["amount"]), d["note"]], 1):
                    cell = ws.cell(i, j, v)
                    if j == 5:
                        cell.number_format = "#,##0"
                i += 1

    # --- Perlu Konfirmasi ---
    ws = wb.create_sheet("Perlu Konfirmasi")
    _hdr(ws, ["No Invoice", "PT", "Toko", "Keterangan Filia", "Nilai (Rp)", "Alasan"],
         [13, 26, 18, 34, 16, 44])
    i = 2
    for h in hasil:
        if h.konfirmasi > 0:
            for j, v in enumerate([h.claim.invoice, h.claim.pt, h.claim.toko,
                                   h.claim.filia_label, round(h.konfirmasi), h.alasan], 1):
                cell = ws.cell(i, j, v)
                if j == 5:
                    cell.number_format = "#,##0"
            i += 1

    # --- Cek Klaim Dobel ---
    ws = wb.create_sheet("Cek Klaim Dobel")
    _hdr(ws, ["Jenis", "PT", "Store", "Produk", "Klaim A", "Klaim B"],
         [15, 24, 10, 34, 34, 34])
    for i, d in enumerate(data["dobel"], 2):
        for j, v in enumerate([d["jenis"], d["pt"], d["store"], d["produk"],
                               d["klaim_a"], d["klaim_b"]], 1):
            ws.cell(i, j, v)

    return wb


def bangun_excel_bytes(data: dict, decisions: dict, batch_label: str) -> bytes:
    wb = bangun_workbook(data, decisions, batch_label)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
